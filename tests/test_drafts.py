"""Draft ingestion: every format's happy path + every user-facing failure
reason from the spec's draft error taxonomy. No network, no DB."""

from __future__ import annotations

from io import BytesIO

import pytest

from research_assistant.ingest.drafts import (
    MAX_DRAFT_CHARS,
    DraftError,
    extract_draft_text,
)


def test_txt_happy_path():
    text, truncated = extract_draft_text("draft.txt", b"hello draft")
    assert text == "hello draft"
    assert truncated is False


def test_md_cp1251_fallback_decoding():
    text, _ = extract_draft_text("draft.md", "черновик".encode("cp1251"))
    assert "черновик" in text


def test_unsupported_extension():
    with pytest.raises(DraftError, match=r"unsupported draft format \(\.rtf\)"):
        extract_draft_text("draft.rtf", b"x")


def test_no_extension():
    with pytest.raises(DraftError, match="unsupported draft format"):
        extract_draft_text("draft", b"x")


def test_file_too_large():
    with pytest.raises(DraftError, match="file too large"):
        extract_draft_text("d.txt", b"x" * (10 * 1024 * 1024 + 1))


def test_empty_draft():
    with pytest.raises(DraftError, match="draft contains no text"):
        extract_draft_text("d.txt", b"   \n  ")


def test_truncation_flag():
    text, truncated = extract_draft_text("d.txt", b"a" * (MAX_DRAFT_CHARS + 10))
    assert truncated is True
    assert len(text) == MAX_DRAFT_CHARS


def test_docx_happy_path(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Draft body paragraph.")
    path = tmp_path / "d.docx"
    doc.save(str(path))
    text, _ = extract_draft_text("d.docx", path.read_bytes())
    assert "Draft body paragraph." in text


def test_docx_corrupt():
    with pytest.raises(DraftError, match="could not read the DOCX file"):
        extract_draft_text("d.docx", b"not a zip archive")


def test_pdf_happy_path():
    fpdf = pytest.importorskip("fpdf")  # fpdf2 lives in the export extra
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(text="Draft pdf text")
    text, _ = extract_draft_text("d.pdf", bytes(pdf.output()))
    assert "Draft pdf text" in text


def test_pdf_corrupt():
    with pytest.raises(DraftError, match="could not read the PDF file"):
        extract_draft_text("d.pdf", b"not a pdf")


def test_pdf_without_text_layer():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)
    with pytest.raises(DraftError, match="no extractable text"):
        extract_draft_text("d.pdf", buf.getvalue())


def test_pdf_encrypted():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("pw")
    buf = BytesIO()
    writer.write(buf)
    with pytest.raises(DraftError, match="password-protected"):
        extract_draft_text("d.pdf", buf.getvalue())
